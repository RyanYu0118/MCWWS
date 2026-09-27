"""把 git 提交历史整理成可导入金山文档多维表格的 xlsx（项目表 / 节点表 / 未归类）。

用法：python tools/scripts/export-project-tracker.py [输出路径]
"""
import re
import subprocess
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO.parent / "流浪世界项目进度.xlsx"
COMMIT_URL = "https://github.com/RyanYu0118/MCWWS/commit/"

PROJECTS = [
    {
        "name": "RS Cursor 安卓端开发",
        "category": "工具",
        "status": "开发中",
        "code": "（不在本仓库）",
        "version": "",
        "scopes": [],
        "keywords": [],
        "note": "独立仓库，节点需手动补录",
    },
    {
        "name": "服务端协同插件研发",
        "category": "服务端插件",
        "status": "开发中",
        "code": "tools/mcwws-world-sync",
        "version": "MCWWS_WorldSync 1.2.9",
        "scopes": ["worldsync"],
        "keywords": ["worldsync", "世界同步"],
        "note": "本机与公网 26.2 世界互斥锁 + 对象存储中转同步",
    },
    {
        "name": "AI 建筑师研发",
        "category": "工具",
        "status": "开发中",
        "code": "tools/mcwws-build-bridge、tools/mcwws-build-bridge-mcp",
        "version": "MCWWS_BuildBridge 1.1.0",
        "scopes": ["build-bridge", "build"],
        "keywords": ["buildbridge"],
        "note": "MCP 直写方块 / FAWE / Litematica 参考库 + 建造 Skill",
    },
    {
        "name": "Axiom 扩展研发",
        "category": "服务端插件 + 客户端模组",
        "status": "开发中",
        "code": "tools/mcwws-axiom-survival、tools/mcwws-axiom-survival-client",
        "version": "MCWWS_AxiomSurvival 1.1.12 / Client 1.4.8",
        "scopes": ["axiom", "axiom-survival", "axiom-survival-client", "axiom-client"],
        "keywords": [r"\baxiom\b"],
        "note": "Axiom 生存扣费、钢笔图层、投影拖入粘贴",
    },
    {
        "name": "地图服务研发",
        "category": "网页与地图",
        "status": "已上线",
        "code": "BlueMap + GIS 网页",
        "version": "",
        "scopes": ["GIS", "map", "bluemap"],
        "keywords": ["地图", r"\bbluemap\b", r"\bgis\b"],
        "note": "BlueMap 网页地图、GIS 建筑标注与编辑",
    },
    {
        "name": "客户端服务开发",
        "category": "网页与地图",
        "status": "已上线",
        "code": "tools/mcwws-web-host、tools/mcwws-web-android",
        "version": "",
        "scopes": ["web", "web-android", "webhost", "网页"],
        "keywords": [r"\bpwa\b", r"\bapk\b", "网页"],
        "note": "网页商城 / PWA / Android WebView 客户端",
    },
    {
        "name": "服务器时延优化",
        "category": "运维",
        "status": "已上线",
        "code": "服务器配置 / 插件升级",
        "version": "",
        "scopes": ["perf"],
        "keywords": ["卡顿", "时延", "卡死", "卡住", "性能", r"\btps\b", r"\bstalls?\b", r"\blag\b", r"\blatency\b"],
        "note": "进服卡顿、TPS、插件拖慢排查",
    },
]

NOISE_SCOPES = {"runtime", "运行时", "用户数据", "userdata", "usercache", "市场状态", "数据", "plugins"}
NOISE_SUBJECT = re.compile(r"同步.{0,12}(运行时|物价|玩家数据)|更新(用户|玩家)数据|运行时状态")
HEAD = re.compile(r"^\ufeff?\s*([^(:\s]+)\s*(?:\(([^)]+)\))?\s*[:：]\s*(.*)$")
VERSION = re.compile(r"\b(\d+\.\d+\.\d+)\b")
TYPE_MAP = {
    "新功能": "功能完成", "feat": "功能完成",
    "修复": "修复", "fix": "修复",
    "性能": "性能优化", "perf": "性能优化",
    "重构": "重构", "refactor": "重构",
    "文档": "文档", "docs": "文档",
    "构建": "构建", "build": "构建",
    "样式": "调整", "格式": "调整", "style": "调整",
    "杂项": "杂项", "chore": "杂项",
}


def git_log():
    raw = subprocess.run(
        ["git", "-c", "i18n.logOutputEncoding=UTF-8", "log", "--reverse",
         "--pretty=format:%H\x1f%ad\x1f%s", "--date=format:%Y-%m-%d %H:%M"],
        cwd=REPO, capture_output=True, check=True,
    ).stdout.decode("utf-8", "replace")
    for line in raw.splitlines():
        full, when, subject = line.split("\x1f", 2)
        yield full, when, subject.strip()


def split_subject(subject):
    """中文标题与英文标题挤在同一行时，只保留中文那半。"""
    m = HEAD.match(subject)
    if not m:
        return "", "", subject
    ctype, scope, rest = m.group(1), m.group(2) or "", m.group(3)
    cut = re.search(r"\s(feat|fix|chore|docs|perf|refactor|style|build|test)\s*\(", rest)
    if cut:
        rest = rest[: cut.start()]
    return ctype, scope, rest.strip()


def classify(scope, text):
    low = f"{scope} {text}".lower()
    for p in PROJECTS:
        if scope in p["scopes"]:
            return p["name"]
    for p in PROJECTS:
        if any(re.search(k, low) for k in p["keywords"]):
            return p["name"]
    return None


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    nodes, others = [], []
    for full, when, subject in git_log():
        if subject.startswith("Merge "):
            continue
        ctype, scope, text = split_subject(subject)
        owned = any(scope in p["scopes"] for p in PROJECTS)
        if scope in NOISE_SCOPES or (not owned and NOISE_SUBJECT.search(text)):
            continue
        project = classify(scope, text)
        ver = VERSION.search(text)
        row = {
            "text": text, "project": project, "type": TYPE_MAP.get(ctype, ctype or "其他"),
            "day": when[:10], "when": when, "scope": scope,
            "version": ver.group(1) if ver else "", "hash": full,
        }
        (nodes if project else others).append(row)

    by_proj = defaultdict(list)
    for n in nodes:
        by_proj[n["project"]].append(n)

    wb = Workbook()
    ws = wb.active
    ws.title = "项目表"
    ws.append(["项目名称", "类别", "当前状态", "负责人", "代码位置", "当前版本",
               "开始日期", "最近更新", "节点数", "进度", "说明"])
    for p in PROJECTS:
        ns = by_proj[p["name"]]
        ws.append([
            p["name"], p["category"], p["status"], "Ryan", p["code"], p["version"],
            date.fromisoformat(ns[0]["day"]) if ns else None,
            date.fromisoformat(ns[-1]["day"]) if ns else None,
            len(ns), None, p["note"],
        ])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    autosize(ws, [22, 20, 10, 8, 44, 34, 12, 12, 8, 8, 44])

    wn = wb.create_sheet("节点表")
    wn.append(["节点名称", "所属项目", "节点类型", "计划日期", "实际日期", "状态",
               "版本号", "提交时间", "Git 提交", "Scope", "Wiki/BookNews", "备注"])
    for n in sorted(nodes, key=lambda x: (x["project"], x["when"])):
        d = date.fromisoformat(n["day"])
        wn.append([n["text"], n["project"], n["type"], d, d, "已完成", n["version"],
                   n["when"], COMMIT_URL + n["hash"], n["scope"], "", ""])
        link = wn.cell(row=wn.max_row, column=9)
        link.hyperlink = link.value
        link.style = "Hyperlink"
    for row in wn.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    autosize(wn, [60, 22, 10, 12, 12, 8, 10, 17, 30, 18, 12, 20])

    wo = wb.create_sheet("未归类提交")
    wo.append(["提交说明", "类型", "Scope", "提交时间", "Git 提交", "拟归入项目"])
    for n in others:
        wo.append([n["text"], n["type"], n["scope"], n["when"], COMMIT_URL + n["hash"], ""])
    autosize(wo, [60, 10, 18, 17, 30, 22])

    wb.save(OUT)
    print(f"{OUT}")
    for p in PROJECTS:
        print(f"  {p['name']}: {len(by_proj[p['name']])}")
    print(f"  未归类: {len(others)}")


if __name__ == "__main__":
    main()
